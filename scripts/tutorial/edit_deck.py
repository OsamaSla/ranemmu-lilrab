# -*- coding: utf-8 -*-
"""Surgical edits to the tutorial deck.

Input : tutorial/1.pptx            (original, never modified)
Output: tutorial/1-edited.pptx     (title slide + closing slide, fixes applied)

Fixes:
  * slide 6  - remove the cropped duplicate screenshot + its connector
  * slide 10 - correct the "إعدادت" typo to "إعدادات" (twice)
  * prepend a title slide and append a closing slide
"""
import copy
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(ROOT, "tutorial", "1.pptx")
DST = os.path.join(ROOT, "tutorial", "1-edited.pptx")
ICON = os.path.join(ROOT, "assets", "images", "icon.png")

BG = RGBColor(0x10, 0x14, 0x18)
FG = RGBColor(0xF2, 0xF4, 0xF5)
ACCENT = RGBColor(0x2A, 0xA9, 0xA2)
MUTED = RGBColor(0x9A, 0xA6, 0xAD)

prs = Presentation(SRC)

# ---------------------------------------------------------------- 1) slide 6
s6 = prs.slides[5]
drop = {"Picture 41", "Straight Arrow Connector 42"}
for sh in list(s6.shapes):
    if sh.name in drop:
        sh._element.getparent().remove(sh._element)

# ---------------------------------------------------------------- 2) typos
fixed = 0
for slide in prs.slides:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if "إعدادت" in r.text:
                    r.text = r.text.replace("إعدادت", "إعدادات")
                    fixed += 1

# ---------------------------------------------------------------- helpers
def rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1")

def blank_layout():
    for lay in prs.slide_layouts:
        if lay.name.lower() == "blank":
            return lay
    return prs.slide_layouts[6]

def add_tb(slide, text, left, top, width, height, size, color,
           bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font="Cairo"):
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        rtl(p)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return box

def new_slide():
    s = prs.slides.add_slide(blank_layout())
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s

SW, SH = prs.slide_width, prs.slide_height
CX = SW // 2

# ---------------------------------------------------------------- title slide
title = new_slide()
if os.path.exists(ICON):
    icon_size = int(SH * 0.34)
    title.shapes.add_picture(ICON, Emu(CX - icon_size // 2), Emu(int(SH * 0.12)),
                             width=Emu(icon_size), height=Emu(icon_size))
add_tb(title, "دليل استخدام تطبيق\nرنموا للرب",
       0, int(SH * 0.52), SW, int(SH * 0.22), 54, FG, bold=True)
add_tb(title, "شرح مبسّط لأهم ميزات البرنامج",
       0, int(SH * 0.75), SW, int(SH * 0.10), 28, ACCENT)
add_tb(title, "إعداد: فريق رنموا للرب",
       0, int(SH * 0.87), SW, int(SH * 0.08), 18, MUTED)

# ---------------------------------------------------------------- closing slide
close = new_slide()
if os.path.exists(ICON):
    icon_size = int(SH * 0.28)
    close.shapes.add_picture(ICON, Emu(CX - icon_size // 2), Emu(int(SH * 0.16)),
                             width=Emu(icon_size), height=Emu(icon_size))
add_tb(close, "شكرًا لمشاهدتكم", 0, int(SH * 0.52), SW, int(SH * 0.14), 48, FG, bold=True)
add_tb(close, "حمّل التطبيق وابدأ التسبيح الآن", 0, int(SH * 0.68), SW, int(SH * 0.10), 28, ACCENT)
add_tb(close, "رنموا للرب — ترنيمات روحية", 0, int(SH * 0.80), SW, int(SH * 0.08), 18, MUTED)

# ---------------------------------------------------------------- move title to front
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
sldIdLst.remove(ids[-2])          # title (added first of the two)
sldIdLst.insert(0, ids[-2])

prs.save(DST)
print(f"saved {os.path.basename(DST)}")
print(f"slides: {len(prs.slides)}   typo runs fixed: {fixed}")
